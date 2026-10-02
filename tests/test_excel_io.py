from pathlib import Path

import openpyxl

import checkpoint as ckpt
import excel_io
from validators import CfRecord


def build_scenario():
    records = [
        CfRecord("RSSMRA85M01H501Z", "RSSMRA85M01H501Z", True),
        CfRecord("VRDLGI90A41F205X", "VRDLGI90A41F205X", True),
        CfRecord("INVALIDCF12345", "INVALIDCF12345", False),
        CfRecord("BNCGPP70C15L219K", "BNCGPP70C15L219K", True),  # non ancora processato
    ]
    entries = {
        "RSSMRA85M01H501Z": ckpt.CheckpointEntry(
            codice_fiscale="RSSMRA85M01H501Z",
            timestamp="2026-07-31T10:00:00",
            esito=ckpt.OUTCOME_YES,
            http_status=200,
            sender_allowed=True,
            preferred_languages=["it_IT"],
            error_detail=None,
            attempts=1,
        ),
        "VRDLGI90A41F205X": ckpt.CheckpointEntry(
            codice_fiscale="VRDLGI90A41F205X",
            timestamp="2026-07-31T10:00:01",
            esito=ckpt.OUTCOME_NO,
            http_status=404,
            sender_allowed=None,
            preferred_languages=None,
            error_detail=None,
            attempts=1,
        ),
    }
    stats = {"total_rows": 5, "blank": 1, "duplicates": 0, "invalid_format": 1, "unique_valid": 3}
    return records, entries, stats


def test_write_report_produces_two_sheets(tmp_path: Path):
    records, entries, stats = build_scenario()
    output_path = tmp_path / "report.xlsx"

    excel_io.write_report(output_path, records, entries, stats, Path("dummy_input.csv"))

    wb = openpyxl.load_workbook(output_path)
    assert wb.sheetnames == ["Riepilogo", "Dettaglio"]


def test_detail_sheet_maps_outcomes_correctly(tmp_path: Path):
    records, entries, stats = build_scenario()
    output_path = tmp_path / "report.xlsx"
    excel_io.write_report(output_path, records, entries, stats, Path("dummy_input.csv"))

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Dettaglio"]
    header = [c.value for c in ws[1]]
    rows = {row[header.index("codice_fiscale")]: row for row in ws.iter_rows(min_row=2, values_only=True)}

    assert rows["RSSMRA85M01H501Z"][header.index("esito")] == ckpt.OUTCOME_YES
    assert rows["VRDLGI90A41F205X"][header.index("esito")] == ckpt.OUTCOME_NO
    assert rows["INVALIDCF12345"][header.index("esito")] == ckpt.OUTCOME_INVALID_FORMAT
    assert rows["BNCGPP70C15L219K"][header.index("esito")] == ckpt.OUTCOME_PENDING
    assert ws.max_row == 5  # header + 4 CF


def test_summary_sheet_has_correct_counts(tmp_path: Path):
    records, entries, stats = build_scenario()
    output_path = tmp_path / "report.xlsx"
    excel_io.write_report(output_path, records, entries, stats, Path("dummy_input.csv"))

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Riepilogo"]
    header = [c.value for c in ws[1]]
    voce_idx, valore_idx = header.index("voce"), header.index("valore")
    counts = {row[voce_idx]: row[valore_idx] for row in ws.iter_rows(min_row=2, values_only=True)}

    assert counts[ckpt.OUTCOME_YES] == 1
    assert counts[ckpt.OUTCOME_NO] == 1
    assert counts[ckpt.OUTCOME_INVALID_FORMAT] == 1
    assert counts[ckpt.OUTCOME_PENDING] == 1
    assert counts["CF univoci considerati"] == 4
